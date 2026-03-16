import os
import asyncio
import json
import re
import uuid
import logging
import tempfile
import shutil
import subprocess
from pathlib import Path
from datetime import datetime, timezone
from typing import List, Optional, Dict, Any

import httpx
from fastapi import FastAPI, APIRouter, UploadFile, File, Form, HTTPException, BackgroundTasks, Request
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from dotenv import load_dotenv
from motor.motor_asyncio import AsyncIOMotorClient
from bson import ObjectId
from docx import Document
import pdfplumber
import mammoth
from comment_docx_parser import parse_docx_comments
from comment_docx_editor import apply_edits_and_save

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')
FRONTEND_BUILD_DIR = ROOT_DIR.parent / 'frontend' / 'build'
FRONTEND_STATIC_DIR = FRONTEND_BUILD_DIR / 'static'

mongo_client = AsyncIOMotorClient(os.environ['MONGO_URL'])
db = mongo_client[os.environ['DB_NAME']]

app = FastAPI(title="Document Filling Agent")
api_router = APIRouter(prefix="/api")

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

STORAGE_DIR = Path(tempfile.gettempdir()) / "docfiller"
STORAGE_DIR.mkdir(exist_ok=True, parents=True)

# In-memory job cache
jobs: Dict[str, Any] = {}
comment_sessions: Dict[str, Any] = {}


class ProjectCreate(BaseModel):
    name: str = Field(..., min_length=1, max_length=120)
    template_id: str
    source_id: str
    ollama_url: Optional[str] = None
    model: Optional[str] = None


class ProjectStartRequest(BaseModel):
    ollama_url: Optional[str] = None
    model: Optional[str] = None


class SettingsPayload(BaseModel):
    ollama_url: str = "http://localhost:11434"
    default_model: str = "llama3.2:3b"


class ProjectApprovePayload(BaseModel):
    replacements: Dict[str, str] = Field(default_factory=dict)


class CommentStudioProcessRequest(BaseModel):
    session_id: str
    model: str
    comment_ids: List[str] | None = None
    ollama_url: str = "http://localhost:11434"


class CommentStudioEditOverride(BaseModel):
    comment_id: str
    new_text: str
    replace_scope: str | None = None


class CommentStudioApplyRequest(BaseModel):
    session_id: str
    edits: List[CommentStudioEditOverride]


class LoginPayload(BaseModel):
    username: str
    password: str


def parse_object_id(id_value: str) -> ObjectId:
    try:
        return ObjectId(id_value)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Invalid id") from exc


def to_public_doc(doc: Dict[str, Any]) -> Dict[str, Any]:
    pub = dict(doc)
    if '_id' in pub:
        pub['_id'] = str(pub['_id'])
    return pub


def apply_replacements_to_text(text: str, replacements: Dict[str, str]) -> str:
    result = text
    for marker, value in (replacements or {}).items():
        result = result.replace(marker, str(value))
    return result

# Marker regex patterns
PLACEHOLDER_PATTERN = re.compile(r'\{\$([^}]+)\}')   # {$name}
RULE_PATTERN = re.compile(r'\[\$([^\]]+)\]')           # [$rule instruction]


# ─── Text Extraction ──────────────────────────────────────────────────────────

def extract_text_from_pdf(path: str) -> str:
    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                pages.append(t)
    return "\n\n".join(pages)


def extract_text_from_docx(path: str) -> str:
    with open(path, "rb") as f:
        result = mammoth.extract_raw_text(f)
    return result.value


def extract_text_from_doc(path: str) -> str:
    tmp = tempfile.mkdtemp()
    try:
        subprocess.run(
            ['libreoffice', '--headless', '--convert-to', 'docx', '--outdir', tmp, path],
            timeout=30, check=True, capture_output=True
        )
        converted = Path(tmp) / f"{Path(path).stem}.docx"
        if converted.exists():
            return extract_text_from_docx(str(converted))
    except Exception as e:
        logger.warning(f"DOC conversion failed: {e}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return ""


def extract_source_text(path: str, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == '.pdf':
        return extract_text_from_pdf(path)
    elif ext == '.docx':
        return extract_text_from_docx(path)
    elif ext == '.doc':
        return extract_text_from_doc(path)
    raise ValueError(f"Unsupported format: {ext}")


# ─── Template Scanning ────────────────────────────────────────────────────────

def scan_text_for_markers(text: str) -> List[Dict]:
    found = []
    for m in PLACEHOLDER_PATTERN.finditer(text):
        found.append({
            "type": "placeholder",
            "marker": m.group(0),
            "name": m.group(1),
        })
    for m in RULE_PATTERN.finditer(text):
        found.append({
            "type": "rule",
            "marker": m.group(0),
            "instruction": m.group(1),
        })
    return found


def scan_docx_for_markers(path: str) -> List[Dict]:
    doc = Document(path)
    seen: set = set()
    results: List[Dict] = []

    def process_para(para):
        text = para.text
        if not text.strip():
            return
        for m in scan_text_for_markers(text):
            if m['marker'] not in seen:
                seen.add(m['marker'])
                results.append(m)

    for para in doc.paragraphs:
        process_para(para)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    process_para(para)
    for section in doc.sections:
        for hf in [section.header, section.footer]:
            if hf:
                for para in hf.paragraphs:
                    process_para(para)
    return results


# ─── Template Filling ─────────────────────────────────────────────────────────

def replace_in_para(para, replacements: Dict[str, str]):
    full = para.text
    if not any(k in full for k in replacements):
        return

    # Single-run replacements first (preserves formatting)
    for run in para.runs:
        if not run.text:
            continue
        new_text = run.text
        for marker, value in replacements.items():
            if marker in new_text:
                new_text = new_text.replace(marker, str(value))
        if new_text != run.text:
            run.text = new_text

    # If markers still span multiple runs, merge into first run
    remaining = para.text
    if any(k in remaining for k in replacements):
        new_full = remaining
        for marker, value in replacements.items():
            new_full = new_full.replace(marker, str(value))
        if para.runs:
            para.runs[0].text = new_full
            for run in para.runs[1:]:
                run.text = ""


def fill_docx_template(template_path: str, output_path: str, replacements: Dict[str, str]):
    doc = Document(template_path)
    for para in doc.paragraphs:
        replace_in_para(para, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_in_para(para, replacements)
    for section in doc.sections:
        for hf in [section.header, section.footer]:
            if hf:
                for para in hf.paragraphs:
                    replace_in_para(para, replacements)
    doc.save(output_path)


def convert_docx_to_pdf(docx_path: str, pdf_path: str) -> bool:
    try:
        import sys
        if sys.platform == 'win32':
            from docx2pdf import convert
            convert(docx_path, pdf_path)
            return Path(pdf_path).exists()
        else:
            out_dir = str(Path(pdf_path).parent)
            subprocess.run(
                ['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', out_dir, docx_path],
                timeout=120, capture_output=True, text=True, check=True
            )
            expected = Path(out_dir) / f"{Path(docx_path).stem}.pdf"
            if expected.exists():
                if str(expected) != pdf_path:
                    shutil.move(str(expected), pdf_path)
                return True
    except Exception as e:
        logger.error(f"PDF conversion error: {e}")
    return False


# ─── Ollama Integration ───────────────────────────────────────────────────────

async def call_ollama(ollama_url: str, model: str, messages: list) -> str:
    async with httpx.AsyncClient(timeout=300.0) as client:
        resp = await client.post(
            f"{ollama_url}/api/chat",
            json={"model": model, "messages": messages, "stream": False,
                  "options": {"temperature": 0.1}}
        )
        resp.raise_for_status()
        return resp.json()['message']['content']


def sanitize_plain_text(text: str) -> str:
    if not text:
        return ""
    out = text.strip()
    if out.startswith("```"):
        out = re.sub(r"^```[^\n]*\n?", "", out)
        out = re.sub(r"\n?```$", "", out)
    out = re.sub(r"(?im)^(Revised text|Final revised full snippet)\s*:\s*", "", out)
    out = re.sub(r"(?m)^#{1,6}\s+", "", out)
    out = re.sub(r"\*\*(.*?)\*\*", r"\1", out)
    out = re.sub(r"__(.*?)__", r"\1", out)
    return out.strip()


async def rewrite_anchor_text(original: str, instruction: str, model: str, ollama_url: str) -> str:
    prompt = (
        "You are a professional document editor working on legal/policy text. "
        "Apply only the requested change to the provided text span. "
        "Return plain text only, no explanation and no markdown.\n\n"
        f"Original text:\n{original}\n\n"
        f"Editor's instruction:\n{instruction}\n\n"
        "Return only revised text:"
    )
    raw = await call_ollama(ollama_url, model, [{"role": "user", "content": prompt}])
    return sanitize_plain_text(raw)


async def rewrite_snippet_with_edit(original_snippet: str, original_span: str, revised_span: str, instruction: str, model: str, ollama_url: str) -> str:
    prompt = (
        "You are a professional policy document editor. "
        "Integrate edited span into the original snippet and return final snippet only.\n\n"
        f"Original full snippet:\n{original_snippet}\n\n"
        f"Original anchored span:\n{original_span}\n\n"
        f"Edited anchored span:\n{revised_span}\n\n"
        f"Instruction:\n{instruction}\n\n"
        "Final revised full snippet:"
    )
    raw = await call_ollama(ollama_url, model, [{"role": "user", "content": prompt}])
    return sanitize_plain_text(raw)


async def rewrite_snippet_with_comments(original_snippet: str, comments: List[Dict[str, str]], model: str, ollama_url: str) -> str:
    comments_block = "\n".join(
        f"- Comment #{c.get('comment_id')}: {c.get('instruction', '')}"
        for c in comments
    )
    prompt = (
        "You are a professional policy document editor. "
        "Apply all reviewer comments to the snippet and return one final snippet only.\n\n"
        f"Original snippet:\n{original_snippet}\n\n"
        f"Comments:\n{comments_block}\n\n"
        "Final revised full snippet:"
    )
    raw = await call_ollama(ollama_url, model, [{"role": "user", "content": prompt}])
    return sanitize_plain_text(raw)


async def call_ollama_streaming(
    ollama_url: str,
    model: str,
    messages: list,
    on_chunk=None,
) -> str:
    chunks: List[str] = []
    async with httpx.AsyncClient(timeout=300.0) as client:
        async with client.stream(
            'POST',
            f"{ollama_url}/api/chat",
            json={
                'model': model,
                'messages': messages,
                'stream': True,
                'options': {'temperature': 0.1}
            }
        ) as resp:
            resp.raise_for_status()
            async for line in resp.aiter_lines():
                if not line:
                    continue
                try:
                    payload = json.loads(line)
                except json.JSONDecodeError:
                    continue
                token = payload.get('message', {}).get('content', '')
                if token:
                    chunks.append(token)
                    if on_chunk is not None:
                        await on_chunk(''.join(chunks))
                if payload.get('done'):
                    break
    return ''.join(chunks).strip()


def extract_json_from_response(text: str) -> dict:
    text = text.strip()
    if text.startswith("```"):
        lines = text.split('\n')
        text = '\n'.join(l for l in lines if not l.startswith("```"))
    # Try to find JSON object
    match = re.search(r'\{[\s\S]*\}', text)
    if match:
        return json.loads(match.group(0))
    return json.loads(text)


async def fill_markers_with_ollama(
    source_text: str,
    markers: List[Dict],
    ollama_url: str,
    model: str,
    job_id: str
) -> Dict[str, str]:
    if not markers:
        return {}

    # Truncate source text to ~12000 chars (~3000 tokens) to fit context
    max_len = 12000
    truncated = source_text[:max_len]
    if len(source_text) > max_len:
        truncated += "\n\n[... document continues ...]"

    # Marker-by-marker streaming allows live typing visualization in the editor page.
    replacements = {}

    if job_id in jobs:
        jobs[job_id]['stream'] = {
            'is_streaming': True,
            'current_marker': None,
            'current_text': '',
            'marker_index': 0,
            'total_markers': len(markers),
        }

    for i, m in enumerate(markers):
        try:
            if m['type'] == 'placeholder':
                prompt = (
                    f"SOURCE DOCUMENT:\n---\n{truncated}\n---\n\n"
                    f"Extract the value for the placeholder '{m['name']}'. "
                    f"Return ONLY the extracted value, nothing else."
                )
            else:
                prompt = (
                    f"SOURCE DOCUMENT:\n---\n{truncated}\n---\n\n"
                    f"Apply this instruction: {m['instruction']}\n"
                    f"Return ONLY the result, nothing else."
                )

            async def on_chunk(partial_text: str):
                if job_id in jobs:
                    jobs[job_id]['stream'] = {
                        'is_streaming': True,
                        'current_marker': m['marker'],
                        'current_text': partial_text,
                        'marker_index': i + 1,
                        'total_markers': len(markers),
                    }

            value = await call_ollama_streaming(
                ollama_url,
                model,
                [
                    {"role": "system", "content": "You are a document analyst. Return only the requested content, no explanation."},
                    {"role": "user", "content": prompt}
                ],
                on_chunk=on_chunk,
            )

            replacements[m['marker']] = value.strip()
        except Exception as e2:
            logger.error(f"Individual fill failed for {m['marker']}: {e2}")
            replacements[m['marker']] = f'[MISSING: insufficient source data for "{m["marker"]}"]'

        # Update progress
        progress = 20 + int((i + 1) / len(markers) * 70)
        if job_id in jobs:
            jobs[job_id].update({
                'progress': progress,
                'markers_filled': i + 1,
                'message': f'Filled {i+1}/{len(markers)} markers...',
                'stream': {
                    'is_streaming': True,
                    'current_marker': m['marker'],
                    'current_text': replacements[m['marker']],
                    'marker_index': i + 1,
                    'total_markers': len(markers),
                }
            })
            await update_job(job_id)

    if job_id in jobs:
        jobs[job_id]['stream'] = {
            'is_streaming': False,
            'current_marker': None,
            'current_text': '',
            'marker_index': len(markers),
            'total_markers': len(markers),
        }

    return replacements


# ─── Job Persistence ──────────────────────────────────────────────────────────

async def update_job(job_id: str):
    if job_id not in jobs:
        return
    data = {k: v for k, v in jobs[job_id].items()
            if k not in ('docx_path', 'pdf_path')}  # Don't store file paths in DB
    if isinstance(data.get('created_at'), datetime):
        data['created_at'] = data['created_at'].isoformat()
    await db.jobs.update_one({'job_id': job_id}, {'$set': data}, upsert=True)


# ─── Background Processing ────────────────────────────────────────────────────

async def process_job(
    job_id: str, source_path: str, source_filename: str,
    template_path: str, template_filename: str,
    ollama_url: str, model: str,
    cleanup_inputs: bool = True
):
    start = datetime.now(timezone.utc)
    try:
        jobs[job_id].update({'status': 'processing', 'message': 'Reading source document...', 'progress': 5})
        await update_job(job_id)

        source_text = extract_source_text(source_path, source_filename)
        if not source_text.strip():
            raise ValueError("Could not extract text from source document")

        jobs[job_id].update({'message': 'Scanning template for markers...', 'progress': 15})
        await update_job(job_id)

        markers = scan_docx_for_markers(template_path)
        jobs[job_id].update({
            'markers': markers,
            'markers_found': len(markers),
            'message': f'Found {len(markers)} marker(s). Filling with AI...',
            'progress': 20
        })
        await update_job(job_id)

        replacements = await fill_markers_with_ollama(
            source_text, markers, ollama_url, model, job_id
        )
        jobs[job_id].update({'replacements': replacements})
        await update_job(job_id)

        jobs[job_id].update({'message': 'Writing filled document...', 'progress': 92})
        await update_job(job_id)

        job_dir = STORAGE_DIR / job_id
        job_dir.mkdir(exist_ok=True)
        docx_out = str(job_dir / "filled_document.docx")
        fill_docx_template(template_path, docx_out, replacements)

        jobs[job_id].update({'message': 'Converting to PDF...', 'progress': 95})
        await update_job(job_id)

        pdf_out = str(job_dir / "filled_document.pdf")
        pdf_ok = convert_docx_to_pdf(docx_out, pdf_out)

        elapsed = (datetime.now(timezone.utc) - start).total_seconds()
        jobs[job_id].update({
            'status': 'done',
            'progress': 100,
            'message': 'Document filled successfully!',
            'markers_filled': len(replacements),
            'processing_time': round(elapsed, 1),
            'has_pdf': pdf_ok,
            'docx_path': docx_out,
            'pdf_path': pdf_out if pdf_ok else None
        })
        await update_job(job_id)

    except Exception as e:
        logger.exception(f"Job {job_id} failed: {e}")
        jobs[job_id].update({'status': 'error', 'message': 'Processing failed', 'error': str(e), 'progress': 0})
        await update_job(job_id)
    finally:
        if cleanup_inputs:
            for p in [source_path, template_path]:
                try:
                    if os.path.exists(p):
                        os.remove(p)
                except Exception:
                    pass


# ─── API Routes ───────────────────────────────────────────────────────────────

@api_router.get("/health")
async def health():
    return {"status": "ok"}


@api_router.post('/auth/login')
async def login(data: LoginPayload):
    configured_user = os.environ.get('APP_LOGIN_USERNAME', 'admin')
    configured_pass = os.environ.get('APP_LOGIN_PASSWORD', 'admin123')

    if data.username != configured_user or data.password != configured_pass:
        raise HTTPException(status_code=401, detail='Invalid username or password')

    return {'ok': True, 'username': configured_user}


@api_router.get("/check-ollama")
async def check_ollama(ollama_url: str = "http://localhost:11434"):
    try:
        async with httpx.AsyncClient(timeout=5.0) as client:
            resp = await client.get(f"{ollama_url}/api/tags")
            if resp.status_code == 200:
                data = resp.json()
                return {"connected": True, "model_count": len(data.get('models', [])), "url": ollama_url}
    except Exception:
        pass
    return {"connected": False, "url": ollama_url}


@api_router.get("/models")
async def list_models(ollama_url: str = "http://localhost:11434"):
    try:
        async with httpx.AsyncClient(timeout=10.0) as client:
            resp = await client.get(f"{ollama_url}/api/tags")
            resp.raise_for_status()
            data = resp.json()
            models = [
                {
                    "name": m.get('name', ''),
                    "size": m.get('size', 0),
                    "parameter_size": m.get('details', {}).get('parameter_size', ''),
                    "family": m.get('details', {}).get('family', ''),
                    "quantization": m.get('details', {}).get('quantization_level', '')
                }
                for m in data.get('models', [])
            ]
            return {"models": models}
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"Cannot connect to Ollama: {str(e)}")


@api_router.post("/analyze-template")
async def analyze_template(template: UploadFile = File(...)):
    if not template.filename.lower().endswith('.docx'):
        raise HTTPException(status_code=400, detail="Template must be a .docx file")
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        tmp.write(await template.read())
        tmp_path = tmp.name
    try:
        markers = scan_docx_for_markers(tmp_path)
        return {
            "filename": template.filename,
            "marker_count": len(markers),
            "placeholders": [m for m in markers if m['type'] == 'placeholder'],
            "rules": [m for m in markers if m['type'] == 'rule']
        }
    finally:
        os.unlink(tmp_path)


@api_router.post("/process")
async def process_documents(
    background_tasks: BackgroundTasks,
    source: UploadFile = File(...),
    template: UploadFile = File(...),
    ollama_url: str = Form("http://localhost:11434"),
    model: str = Form(...)
):
    src_ext = Path(source.filename).suffix.lower()
    if src_ext not in ['.pdf', '.doc', '.docx']:
        raise HTTPException(status_code=400, detail="Source must be PDF, DOC, or DOCX")
    if not template.filename.lower().endswith('.docx'):
        raise HTTPException(status_code=400, detail="Template must be a .docx file")

    job_id = str(uuid.uuid4())
    source_path = str(STORAGE_DIR / f"{job_id}_source{src_ext}")
    template_path = str(STORAGE_DIR / f"{job_id}_template.docx")

    with open(source_path, 'wb') as f:
        f.write(await source.read())
    with open(template_path, 'wb') as f:
        f.write(await template.read())

    jobs[job_id] = {
        'job_id': job_id,
        'status': 'pending',
        'progress': 0,
        'message': 'Job queued...',
        'markers_found': 0,
        'markers_filled': 0,
        'source_filename': source.filename,
        'template_filename': template.filename,
        'processing_time': None,
        'error': None,
        'created_at': datetime.now(timezone.utc)
    }
    await update_job(job_id)

    background_tasks.add_task(
        process_job, job_id=job_id,
        source_path=source_path, source_filename=source.filename,
        template_path=template_path, template_filename=template.filename,
        ollama_url=ollama_url, model=model
    )
    return {"job_id": job_id, "status": "pending"}


@api_router.get("/jobs/{job_id}")
async def get_job(job_id: str):
    if job_id in jobs:
        job = jobs[job_id]
        data = {k: v for k, v in job.items() if k not in ('docx_path', 'pdf_path')}
        if isinstance(data.get('created_at'), datetime):
            data['created_at'] = data['created_at'].isoformat()
        job_dir = STORAGE_DIR / job_id
        data['has_docx'] = (job_dir / "filled_document.docx").exists()
        data['has_pdf'] = (job_dir / "filled_document.pdf").exists()
        return data
    job = await db.jobs.find_one({'job_id': job_id}, {'_id': 0})
    if job:
        job_dir = STORAGE_DIR / job_id
        job['has_docx'] = (job_dir / "filled_document.docx").exists()
        job['has_pdf'] = (job_dir / "filled_document.pdf").exists()
        return job
    raise HTTPException(status_code=404, detail="Job not found")


@api_router.get("/download/{job_id}/{fmt}")
async def download_file(job_id: str, fmt: str):
    if fmt not in ('docx', 'pdf'):
        raise HTTPException(status_code=400, detail="Format must be 'docx' or 'pdf'")

    # Check in-memory cache first, then MongoDB
    job = jobs.get(job_id)
    if not job:
        job = await db.jobs.find_one({'job_id': job_id}, {'_id': 0})
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    if job.get('status') != 'done':
        raise HTTPException(status_code=400, detail="Job not completed")

    # Derive file paths from job_id (predictable pattern, survives server restart)
    job_dir = STORAGE_DIR / job_id
    if fmt == 'docx':
        fpath = str(job_dir / "filled_document.docx")
        media = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        fname = 'filled_document.docx'
    else:
        fpath = str(job_dir / "filled_document.pdf")
        media = 'application/pdf'
        fname = 'filled_document.pdf'

    if not os.path.exists(fpath):
        raise HTTPException(status_code=404, detail=f"{fmt.upper()} file not available")

    return FileResponse(path=fpath, media_type=media, filename=fname,
                        headers={"Content-Disposition": f"attachment; filename={fname}"})


@api_router.post('/comment-studio/upload')
async def upload_comment_studio_docx(file: UploadFile = File(...)):
    if not (file.filename or '').lower().endswith('.docx'):
        raise HTTPException(status_code=400, detail='Only .docx files are supported')

    session_id = str(uuid.uuid4())
    tmp_dir = tempfile.mkdtemp()
    upload_path = os.path.join(tmp_dir, f'{session_id}_input.docx')

    with open(upload_path, 'wb') as f:
        f.write(await file.read())

    try:
        result = parse_docx_comments(upload_path)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f'Failed to parse DOCX: {e}')

    comment_sessions[session_id] = {
        'parse_result': result,
        'upload_path': upload_path,
        'tmp_dir': tmp_dir,
    }

    comments = [
        {
            'comment_id': c.comment_id,
            'author': c.comment_author,
            'comment_date': c.comment_date,
            'comment_text': c.comment_text,
            'anchored_text': c.anchored_text,
            'source_part': c.source_part,
            'paragraph_indices': c.paragraph_indices,
        }
        for c in result.comments
    ]

    return {
        'session_id': session_id,
        'filename': file.filename,
        'comment_count': len(comments),
        'comments': comments,
    }


@api_router.post('/comment-studio/process')
async def process_comment_studio(req: CommentStudioProcessRequest):
    session = comment_sessions.get(req.session_id)
    if not session:
        raise HTTPException(status_code=404, detail='Session not found. Re-upload the file.')

    parse_result = session['parse_result']
    anchors = parse_result.comments
    if req.comment_ids:
        selected_ids = set(req.comment_ids)
        anchors = [a for a in anchors if a.comment_id in selected_ids]

    suggestions = []
    errors = []
    infos = []
    para_groups: Dict[tuple, list] = {}

    for a in anchors:
        if a.paragraph_indices:
            para_idx = a.paragraph_indices[0]
            source_part = a.source_part or 'word/document.xml'
            para_groups.setdefault((source_part, para_idx), []).append(a)

    processed_comment_ids: set[str] = set()

    for _, group in para_groups.items():
        if len(group) <= 1:
            continue

        leader = sorted(group, key=lambda a: int(a.comment_id) if str(a.comment_id).isdigit() else 999999)[0]
        snippet = (leader.context_snippet or '').strip()
        valid_group = [g for g in group if (g.comment_text or '').strip()]

        if not snippet:
            for g in group:
                processed_comment_ids.add(g.comment_id)
                errors.append({'comment_id': g.comment_id, 'error': 'Missing paragraph context for merged edit'})
            continue

        try:
            merged_text = await rewrite_snippet_with_comments(
                original_snippet=snippet,
                comments=[{'comment_id': g.comment_id, 'instruction': g.comment_text} for g in valid_group],
                model=req.model,
                ollama_url=req.ollama_url,
            )
            if not merged_text:
                raise ValueError('Model returned empty paragraph rewrite')

            comment_summary = ' | '.join(f"#{g.comment_id}: {g.comment_text}" for g in valid_group)
            suggestions.append({
                'comment_id': leader.comment_id,
                'original_text': snippet,
                'comment': f'Merged paragraph comments: {comment_summary}',
                'new_text': merged_text,
                'replace_scope': 'paragraph',
                'related_comment_ids': [g.comment_id for g in valid_group],
            })
            for g in group:
                processed_comment_ids.add(g.comment_id)
                if g.comment_id != leader.comment_id:
                    infos.append({'comment_id': g.comment_id, 'message': f'Merged into paragraph rewrite led by #{leader.comment_id}'})
        except Exception as e:
            for g in group:
                processed_comment_ids.add(g.comment_id)
                errors.append({'comment_id': g.comment_id, 'error': f'Merged paragraph rewrite failed: {e}'})

    for anchor in anchors:
        if anchor.comment_id in processed_comment_ids:
            continue
        if not (anchor.anchored_text or '').strip():
            errors.append({'comment_id': anchor.comment_id, 'error': 'No anchored text found'})
            continue
        if not (anchor.comment_text or '').strip():
            errors.append({'comment_id': anchor.comment_id, 'error': 'Comment is empty'})
            continue

        try:
            anchor_edit = await rewrite_anchor_text(
                original=anchor.anchored_text,
                instruction=anchor.comment_text,
                model=req.model,
                ollama_url=req.ollama_url,
            )

            new_text = anchor_edit
            replace_scope = 'anchor'

            if anchor.paragraph_indices and anchor.context_snippet:
                merged = await rewrite_snippet_with_edit(
                    original_snippet=anchor.context_snippet,
                    original_span=anchor.anchored_text,
                    revised_span=anchor_edit,
                    instruction=anchor.comment_text,
                    model=req.model,
                    ollama_url=req.ollama_url,
                )
                if merged:
                    new_text = merged
                    replace_scope = 'paragraph'

            suggestions.append({
                'comment_id': anchor.comment_id,
                'original_text': anchor.anchored_text,
                'comment': anchor.comment_text,
                'new_text': new_text,
                'replace_scope': replace_scope,
                'related_comment_ids': [anchor.comment_id],
            })
        except Exception as e:
            errors.append({'comment_id': anchor.comment_id, 'error': str(e)})

    session['suggestions'] = {s['comment_id']: s for s in suggestions}

    return {'suggestions': suggestions, 'errors': errors, 'infos': infos}


@api_router.post('/comment-studio/apply')
async def apply_comment_studio(req: CommentStudioApplyRequest):
    session = comment_sessions.get(req.session_id)
    if not session:
        raise HTTPException(status_code=404, detail='Session not found')

    parse_result = session['parse_result']
    tmp_dir = session['tmp_dir']
    out_path = os.path.join(tmp_dir, f'{req.session_id}_output.docx')
    cached_suggestions = session.get('suggestions', {})

    edits = [
        {
            'comment_id': e.comment_id,
            'new_text': e.new_text,
            'replace_scope': e.replace_scope or 'anchor',
        }
        for e in req.edits
    ]

    resolve_ids: set[str] = set()
    for e in req.edits:
        resolve_ids.add(e.comment_id)
        s = cached_suggestions.get(e.comment_id)
        if s and isinstance(s.get('related_comment_ids'), list):
            for cid in s['related_comment_ids']:
                resolve_ids.add(str(cid))

    try:
        apply_edits_and_save(
            all_files=parse_result.all_files,
            edits=edits,
            resolve_comment_ids=list(resolve_ids),
            output_path=out_path,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f'Failed to apply edits: {e}')

    session['output_path'] = out_path
    return {'download_ready': True, 'session_id': req.session_id}


@api_router.get('/comment-studio/download/{session_id}')
async def download_comment_studio_output(session_id: str):
    session = comment_sessions.get(session_id)
    if not session or 'output_path' not in session:
        raise HTTPException(status_code=404, detail='No output file found for this session.')
    return FileResponse(
        session['output_path'],
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        filename='edited_document.docx',
    )


@api_router.get('/dashboard/stats')
async def get_dashboard_stats():
    total_jobs = await db.jobs.count_documents({})
    running_jobs = await db.jobs.count_documents({'status': {'$in': ['pending', 'processing']}})
    return {
        'templates': await db.templates.count_documents({}),
        'sources': await db.sources.count_documents({}),
        'projects': await db.projects.count_documents({}),
        'jobs': total_jobs,
        'running_jobs': running_jobs,
        'users': 1,
    }


@api_router.get('/templates')
async def get_templates():
    docs = []
    async for t in db.templates.find().sort('created_at', -1):
        docs.append(to_public_doc(t))
    return docs


@api_router.post('/templates')
async def create_template(name: str = Form(...), file: UploadFile = File(...)):
    if not file.filename.lower().endswith('.docx'):
        raise HTTPException(status_code=400, detail='Template must be a .docx file')

    templates_dir = STORAGE_DIR / 'library' / 'templates'
    templates_dir.mkdir(parents=True, exist_ok=True)
    ext = Path(file.filename).suffix.lower()
    stored_name = f"{uuid.uuid4()}{ext}"
    file_path = templates_dir / stored_name
    with open(file_path, 'wb') as out_file:
        out_file.write(await file.read())

    res = await db.templates.insert_one({
        'name': name,
        'filename': file.filename,
        'stored_filename': stored_name,
        'path': str(file_path),
        'size': file_path.stat().st_size,
        'created_at': datetime.now(timezone.utc)
    })
    return {'id': str(res.inserted_id), 'name': name}


@api_router.put('/templates/{template_id}')
async def update_template(
    template_id: str,
    name: Optional[str] = Form(None),
    file: Optional[UploadFile] = File(None)
):
    oid = parse_object_id(template_id)
    doc = await db.templates.find_one({'_id': oid})
    if not doc:
        raise HTTPException(status_code=404, detail='Template not found')

    update: Dict[str, Any] = {}
    if name:
        update['name'] = name

    if file is not None:
        if not file.filename.lower().endswith('.docx'):
            raise HTTPException(status_code=400, detail='Template must be a .docx file')
        templates_dir = STORAGE_DIR / 'library' / 'templates'
        templates_dir.mkdir(parents=True, exist_ok=True)
        ext = Path(file.filename).suffix.lower()
        stored_name = f"{uuid.uuid4()}{ext}"
        file_path = templates_dir / stored_name
        with open(file_path, 'wb') as out_file:
            out_file.write(await file.read())
        update.update({
            'filename': file.filename,
            'stored_filename': stored_name,
            'path': str(file_path),
            'size': file_path.stat().st_size
        })
        old_path = doc.get('path')
        if old_path and Path(old_path).exists():
            try:
                os.remove(old_path)
            except Exception:
                logger.warning(f'Could not remove old template file: {old_path}')

    if not update:
        raise HTTPException(status_code=400, detail='Nothing to update')

    await db.templates.update_one({'_id': oid}, {'$set': update})
    return {'status': 'ok'}


@api_router.delete('/templates/{template_id}')
async def delete_template(template_id: str):
    oid = parse_object_id(template_id)
    doc = await db.templates.find_one({'_id': oid})
    if not doc:
        raise HTTPException(status_code=404, detail='Template not found')
    await db.templates.delete_one({'_id': oid})
    file_path = doc.get('path')
    if file_path and Path(file_path).exists():
        try:
            os.remove(file_path)
        except Exception:
            logger.warning(f'Could not remove template file: {file_path}')
    return {'status': 'ok'}


@api_router.get('/sources')
async def get_sources():
    docs = []
    async for s in db.sources.find().sort('created_at', -1):
        docs.append(to_public_doc(s))
    return docs


@api_router.post('/sources')
async def create_source(name: str = Form(...), file: UploadFile = File(...)):
    ext = Path(file.filename).suffix.lower()
    if ext not in ['.pdf', '.doc', '.docx']:
        raise HTTPException(status_code=400, detail='Source must be PDF, DOC, or DOCX')

    sources_dir = STORAGE_DIR / 'library' / 'sources'
    sources_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid.uuid4()}{ext}"
    file_path = sources_dir / stored_name
    with open(file_path, 'wb') as out_file:
        out_file.write(await file.read())

    res = await db.sources.insert_one({
        'name': name,
        'filename': file.filename,
        'stored_filename': stored_name,
        'path': str(file_path),
        'size': file_path.stat().st_size,
        'created_at': datetime.now(timezone.utc)
    })
    return {'id': str(res.inserted_id), 'name': name}


@api_router.put('/sources/{source_id}')
async def update_source(
    source_id: str,
    name: Optional[str] = Form(None),
    file: Optional[UploadFile] = File(None)
):
    oid = parse_object_id(source_id)
    doc = await db.sources.find_one({'_id': oid})
    if not doc:
        raise HTTPException(status_code=404, detail='Source not found')

    update: Dict[str, Any] = {}
    if name:
        update['name'] = name

    if file is not None:
        ext = Path(file.filename).suffix.lower()
        if ext not in ['.pdf', '.doc', '.docx']:
            raise HTTPException(status_code=400, detail='Source must be PDF, DOC, or DOCX')
        sources_dir = STORAGE_DIR / 'library' / 'sources'
        sources_dir.mkdir(parents=True, exist_ok=True)
        stored_name = f"{uuid.uuid4()}{ext}"
        file_path = sources_dir / stored_name
        with open(file_path, 'wb') as out_file:
            out_file.write(await file.read())
        update.update({
            'filename': file.filename,
            'stored_filename': stored_name,
            'path': str(file_path),
            'size': file_path.stat().st_size
        })
        old_path = doc.get('path')
        if old_path and Path(old_path).exists():
            try:
                os.remove(old_path)
            except Exception:
                logger.warning(f'Could not remove old source file: {old_path}')

    if not update:
        raise HTTPException(status_code=400, detail='Nothing to update')

    await db.sources.update_one({'_id': oid}, {'$set': update})
    return {'status': 'ok'}


@api_router.delete('/sources/{source_id}')
async def delete_source(source_id: str):
    oid = parse_object_id(source_id)
    doc = await db.sources.find_one({'_id': oid})
    if not doc:
        raise HTTPException(status_code=404, detail='Source not found')
    await db.sources.delete_one({'_id': oid})
    file_path = doc.get('path')
    if file_path and Path(file_path).exists():
        try:
            os.remove(file_path)
        except Exception:
            logger.warning(f'Could not remove source file: {file_path}')
    return {'status': 'ok'}


@api_router.get('/projects')
async def get_projects():
    templates: Dict[str, str] = {}
    async for t in db.templates.find({}, {'name': 1}):
        templates[str(t['_id'])] = t.get('name', 'Unknown template')

    sources: Dict[str, str] = {}
    async for s in db.sources.find({}, {'name': 1}):
        sources[str(s['_id'])] = s.get('name', 'Unknown source')

    docs = []
    async for p in db.projects.find().sort('created_at', -1):
        item = to_public_doc(p)
        item['template_name'] = templates.get(item.get('template_id', ''), 'Unknown template')
        item['source_name'] = sources.get(item.get('source_id', ''), 'Unknown source')
        docs.append(item)
    return docs


@api_router.post('/projects')
async def create_project(data: ProjectCreate):
    template = await db.templates.find_one({'_id': parse_object_id(data.template_id)})
    if not template:
        raise HTTPException(status_code=404, detail='Template not found')
    source = await db.sources.find_one({'_id': parse_object_id(data.source_id)})
    if not source:
        raise HTTPException(status_code=404, detail='Source not found')

    payload = data.model_dump()
    payload.update({
        'status': 'idle',
        'last_job_id': None,
        'created_at': datetime.now(timezone.utc)
    })
    res = await db.projects.insert_one(payload)
    return {'id': str(res.inserted_id)}


@api_router.put('/projects/{project_id}')
async def update_project(project_id: str, data: ProjectCreate):
    project_oid = parse_object_id(project_id)
    template = await db.templates.find_one({'_id': parse_object_id(data.template_id)})
    if not template:
        raise HTTPException(status_code=404, detail='Template not found')
    source = await db.sources.find_one({'_id': parse_object_id(data.source_id)})
    if not source:
        raise HTTPException(status_code=404, detail='Source not found')

    await db.projects.update_one(
        {'_id': project_oid},
        {'$set': {**data.model_dump(), 'updated_at': datetime.now(timezone.utc)}}
    )
    return {'status': 'ok'}


@api_router.delete('/projects/{project_id}')
async def delete_project(project_id: str):
    project_oid = parse_object_id(project_id)
    await db.projects.delete_one({'_id': project_oid})
    return {'status': 'ok'}


@api_router.post('/projects/{project_id}/start')
async def start_project(project_id: str, background_tasks: BackgroundTasks, payload: ProjectStartRequest):
    project_oid = parse_object_id(project_id)
    project = await db.projects.find_one({'_id': project_oid})
    if not project:
        raise HTTPException(status_code=404, detail='Project not found')

    template = await db.templates.find_one({'_id': parse_object_id(project['template_id'])})
    source = await db.sources.find_one({'_id': parse_object_id(project['source_id'])})
    if not template or not source:
        raise HTTPException(status_code=400, detail='Project has missing template/source')

    template_path = template.get('path')
    source_path = source.get('path')
    if not template_path or not source_path or not Path(template_path).exists() or not Path(source_path).exists():
        raise HTTPException(status_code=400, detail='Template or source file is unavailable on disk')

    saved_settings = await db.settings.find_one({}) or {}
    ollama_url = payload.ollama_url or project.get('ollama_url') or saved_settings.get('ollama_url') or 'http://localhost:11434'
    model = payload.model or project.get('model') or saved_settings.get('default_model') or 'llama3.2:3b'

    job_id = str(uuid.uuid4())
    src_ext = Path(source.get('filename', '')).suffix.lower() or '.pdf'
    source_tmp = str(STORAGE_DIR / f"{job_id}_source{src_ext}")
    template_tmp = str(STORAGE_DIR / f"{job_id}_template.docx")
    shutil.copyfile(source_path, source_tmp)
    shutil.copyfile(template_path, template_tmp)

    jobs[job_id] = {
        'job_id': job_id,
        'status': 'pending',
        'progress': 0,
        'message': 'Job queued...',
        'markers_found': 0,
        'markers_filled': 0,
        'source_filename': source.get('filename', 'source.docx'),
        'template_filename': template.get('filename', 'template.docx'),
        'processing_time': None,
        'error': None,
        'project_id': project_id,
        'created_at': datetime.now(timezone.utc)
    }
    await update_job(job_id)

    await db.projects.update_one(
        {'_id': project_oid},
        {'$set': {'status': 'processing', 'last_job_id': job_id, 'updated_at': datetime.now(timezone.utc)}}
    )

    background_tasks.add_task(
        process_job, job_id=job_id,
        source_path=source_tmp, source_filename=source.get('filename', 'source.docx'),
        template_path=template_tmp, template_filename=template.get('filename', 'template.docx'),
        ollama_url=ollama_url, model=model,
        cleanup_inputs=True
    )
    return {'job_id': job_id, 'status': 'pending'}


@api_router.get('/projects/{project_id}/editor/{job_id}')
async def get_project_editor_payload(project_id: str, job_id: str):
    project = await db.projects.find_one({'_id': parse_object_id(project_id)})
    if not project:
        raise HTTPException(status_code=404, detail='Project not found')

    template = await db.templates.find_one({'_id': parse_object_id(project['template_id'])})
    if not template:
        raise HTTPException(status_code=404, detail='Template not found')

    template_path = template.get('path')
    if not template_path or not Path(template_path).exists():
        raise HTTPException(status_code=400, detail='Template file is unavailable on disk')

    template_text = extract_text_from_docx(template_path)

    job = jobs.get(job_id)
    if not job:
        job = await db.jobs.find_one({'job_id': job_id}, {'_id': 0})
    if not job:
        raise HTTPException(status_code=404, detail='Job not found')

    replacements = job.get('approved_replacements') or job.get('replacements') or {}
    markers = job.get('markers') or scan_docx_for_markers(template_path)

    return {
        'project_id': project_id,
        'job_id': job_id,
        'project_name': project.get('name', 'Project'),
        'status': job.get('status', 'pending'),
        'message': job.get('message', ''),
        'progress': job.get('progress', 0),
        'template_name': template.get('name', ''),
        'template_text': template_text,
        'markers': markers,
        'replacements': replacements,
        'preview_text': apply_replacements_to_text(template_text, replacements),
        'stream': job.get('stream', {}),
        'has_docx': (STORAGE_DIR / job_id / 'filled_document.docx').exists(),
        'has_pdf': (STORAGE_DIR / job_id / 'filled_document.pdf').exists(),
        'approved': bool(job.get('approved', False)),
    }


@api_router.get('/projects/{project_id}/editor/{job_id}/stream')
async def stream_project_editor(project_id: str, job_id: str, request: Request):
    async def event_generator():
        while True:
            if await request.is_disconnected():
                break

            job = jobs.get(job_id)
            if not job:
                db_job = await db.jobs.find_one({'job_id': job_id}, {'_id': 0})
                if not db_job:
                    payload = {'error': 'Job not found', 'job_id': job_id}
                    yield f"data: {json.dumps(payload)}\n\n"
                    break
                job = db_job

            payload = {
                'job_id': job_id,
                'status': job.get('status', 'pending'),
                'progress': job.get('progress', 0),
                'message': job.get('message', ''),
                'markers_filled': job.get('markers_filled', 0),
                'markers_found': job.get('markers_found', 0),
                'stream': job.get('stream', {}),
            }
            yield f"data: {json.dumps(payload)}\n\n"

            if payload['status'] in ('done', 'error') and not payload['stream'].get('is_streaming', False):
                break

            await asyncio.sleep(0.45)

    return StreamingResponse(event_generator(), media_type='text/event-stream')


@api_router.post('/projects/{project_id}/editor/{job_id}/approve')
async def approve_project_changes(project_id: str, job_id: str, payload: ProjectApprovePayload):
    project = await db.projects.find_one({'_id': parse_object_id(project_id)})
    if not project:
        raise HTTPException(status_code=404, detail='Project not found')

    template = await db.templates.find_one({'_id': parse_object_id(project['template_id'])})
    if not template:
        raise HTTPException(status_code=404, detail='Template not found')

    template_path = template.get('path')
    if not template_path or not Path(template_path).exists():
        raise HTTPException(status_code=400, detail='Template file is unavailable on disk')

    job_dir = STORAGE_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)

    docx_out = str(job_dir / 'filled_document.docx')
    fill_docx_template(template_path, docx_out, payload.replacements)
    pdf_out = str(job_dir / 'filled_document.pdf')
    has_pdf = convert_docx_to_pdf(docx_out, pdf_out)

    if job_id in jobs:
        jobs[job_id].update({
            'approved': True,
            'approved_replacements': payload.replacements,
            'status': 'done',
            'message': 'Approved and finalized by user',
            'has_pdf': has_pdf,
            'docx_path': docx_out,
            'pdf_path': pdf_out if has_pdf else None
        })
        await update_job(job_id)
    else:
        await db.jobs.update_one(
            {'job_id': job_id},
            {
                '$set': {
                    'approved': True,
                    'approved_replacements': payload.replacements,
                    'status': 'done',
                    'message': 'Approved and finalized by user',
                    'has_pdf': has_pdf,
                }
            },
            upsert=True
        )

    await db.projects.update_one(
        {'_id': parse_object_id(project_id)},
        {'$set': {'status': 'approved', 'last_job_id': job_id, 'updated_at': datetime.now(timezone.utc)}}
    )

    return {'status': 'ok', 'has_docx': True, 'has_pdf': has_pdf}


@api_router.get('/settings')
async def get_settings():
    s = await db.settings.find_one({})
    if not s:
        return {'ollama_url': 'http://localhost:11434', 'default_model': 'llama3.2:3b'}
    return {**s, '_id': str(s['_id'])}


@api_router.post('/settings')
async def save_settings(data: SettingsPayload):
    await db.settings.update_one({}, {'$set': data.model_dump()}, upsert=True)
    return {'status': 'ok'}


app.include_router(api_router)


if FRONTEND_STATIC_DIR.exists():
    app.mount('/static', StaticFiles(directory=str(FRONTEND_STATIC_DIR)), name='frontend-static')


if FRONTEND_BUILD_DIR.exists():
    @app.get('/', include_in_schema=False)
    async def serve_frontend_index():
        return FileResponse(str(FRONTEND_BUILD_DIR / 'index.html'))


    @app.get('/{full_path:path}', include_in_schema=False)
    async def serve_frontend_file_or_spa(full_path: str):
        if full_path.startswith('api/'):
            raise HTTPException(status_code=404, detail='Not found')

        requested = FRONTEND_BUILD_DIR / full_path
        if requested.exists() and requested.is_file():
            return FileResponse(str(requested))

        return FileResponse(str(FRONTEND_BUILD_DIR / 'index.html'))


@app.on_event("shutdown")
async def shutdown():
    mongo_client.close()

