import re
import uuid
import logging
from pathlib import Path
from docx import Document
from typing import List, Dict, Any

logger = logging.getLogger(__name__)

PLACEHOLDER_RE = re.compile(r"\[([^\[\]\n]{1,200})\]")


def classify_placeholder(inner: str) -> str:
    """'simple' = ALL_CAPS_WITH_UNDERSCORES, 'instruction' = sentence-like."""
    words = inner.strip().split()
    if not words:
        return "simple"
    all_upper = all(w.replace("_", "").replace(" ", "").isupper() for w in words if w)
    return "simple" if (all_upper and len(inner) < 60) else "instruction"


class TemplateParser:
    def parse(self, template_path: str) -> Dict[str, Any]:
        doc = Document(template_path)
        template_maps = []
        all_paras = self._all_paragraphs(doc)

        para_per_page = 40
        for entry in all_paras:
            para = entry["paragraph"]
            idx = entry["index"]
            page_num = max(1, idx // para_per_page + 1)

            full_text = "".join(run.text for run in para.runs)
            for match in PLACEHOLDER_RE.finditer(full_text):
                ph_text = match.group(0)
                inner = match.group(1)
                ph_type = classify_placeholder(inner)

                s, e = match.start(), match.end()
                ctx_before = full_text[max(0, s - 80) : s].strip()
                ctx_after = full_text[e : e + 80].strip()

                ph_id = str(uuid.uuid4())
                template_maps.append(
                    {
                        "_id": ph_id,
                        "placeholder_id": ph_id,
                        "placeholder_text": ph_text,
                        "placeholder_type": ph_type,
                        "page_number": page_num,
                        "paragraph_index": idx,
                        "context_before": ctx_before,
                        "context_after": ctx_after,
                    }
                )

        logger.info(f"Found {len(template_maps)} placeholders in template")
        return {"template_maps": template_maps}

    def _all_paragraphs(self, doc: Document) -> List[Dict]:
        """Iterate body elements in document order (paragraphs + table cells)."""
        result = []
        idx = 0

        def _visit(element):
            nonlocal idx
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag
            if tag == "p":
                try:
                    from docx.text.paragraph import Paragraph
                    para = Paragraph(element, doc)
                    result.append({"paragraph": para, "index": idx})
                    idx += 1
                except Exception:
                    pass
            elif tag == "tbl":
                try:
                    from docx.table import Table
                    table = Table(element, doc)
                    for row in table.rows:
                        for cell in row.cells:
                            for para in cell.paragraphs:
                                result.append({"paragraph": para, "index": idx})
                                idx += 1
                except Exception:
                    pass

        for child in doc.element.body:
            _visit(child)

        return result
