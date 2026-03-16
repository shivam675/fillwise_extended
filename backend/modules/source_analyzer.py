import json
import logging
import numpy as np
from pathlib import Path
from typing import List, Dict, Optional

logger = logging.getLogger(__name__)

CHUNK_SIZE = 900  # characters


class SourceAnalyzer:
    def extract_and_chunk(self, source_path: str, source_ext: str) -> List[Dict]:
        if source_ext == ".docx":
            pages = self._extract_docx(source_path)
        elif source_ext == ".pdf":
            pages = self._extract_pdf(source_path)
        else:
            raise ValueError(f"Unsupported format: {source_ext}")

        chunks = []
        for page_num, text in pages:
            chunks.extend(self._chunk_text(text, page_num))

        logger.info(f"Extracted {len(chunks)} chunks from source")
        return chunks

    def _extract_docx(self, path: str) -> List[tuple]:
        from docx import Document
        doc = Document(path)
        pages, current, count, page_num = [], [], 0, 1

        for para in doc.paragraphs:
            t = para.text.strip()
            if t:
                current.append(t)
            count += 1
            if count >= 40:
                pages.append((page_num, "\n".join(current)))
                current, count, page_num = [], 0, page_num + 1

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        current.append(t)

        if current:
            pages.append((page_num, "\n".join(current)))
        return pages

    def _extract_pdf(self, path: str) -> List[tuple]:
        try:
            import fitz
            doc = fitz.open(path)
            pages = []
            for i, page in enumerate(doc):
                text = page.get_text()
                if text.strip():
                    pages.append((i + 1, text))
            doc.close()
            return pages
        except ImportError:
            from PyPDF2 import PdfReader
            reader = PdfReader(path)
            pages = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append((i + 1, text))
            return pages

    def _chunk_text(self, text: str, page_num: int) -> List[Dict]:
        paras = [p.strip() for p in text.split("\n") if p.strip()]
        chunks, current, size = [], [], 0

        for para in paras:
            if size + len(para) > CHUNK_SIZE and current:
                chunks.append({"text": "\n".join(current), "page_number": page_num})
                current = current[-1:]
                size = len(current[0]) if current else 0
            current.append(para)
            size += len(para)

        if current:
            chunks.append({"text": "\n".join(current), "page_number": page_num})
        return chunks

    def build_faiss_index(self, chunk_docs: List[Dict], session_dir: str):
        try:
            import faiss
        except ImportError:
            logger.warning("faiss-cpu not available, skipping index build")
            return

        valid = [(c["_id"], c["embedding_vector"]) for c in chunk_docs if c.get("embedding_vector")]
        if not valid:
            logger.warning("No embeddings available for FAISS index")
            return

        dim = len(valid[0][1])
        ids = [cid for cid, _ in valid]
        embs = np.array([e for _, e in valid], dtype=np.float32)
        faiss.normalize_L2(embs)

        index = faiss.IndexFlatIP(dim)
        index.add(embs)

        idx_path = Path(session_dir) / "faiss.index"
        ids_path = Path(session_dir) / "faiss_ids.json"
        faiss.write_index(index, str(idx_path))
        ids_path.write_text(json.dumps(ids))
        logger.info(f"FAISS index built: {len(valid)} vectors, dim={dim}")

    def search_similar_chunks(
        self, session_dir: str, query_embedding: List[float], top_k: int = 5
    ) -> List[str]:
        try:
            import faiss
            idx_path = Path(session_dir) / "faiss.index"
            ids_path = Path(session_dir) / "faiss_ids.json"
            if not idx_path.exists():
                return []

            index = faiss.read_index(str(idx_path))
            ids = json.loads(ids_path.read_text())

            q = np.array([query_embedding], dtype=np.float32)
            faiss.normalize_L2(q)
            _, indices = index.search(q, min(top_k, len(ids)))

            return [ids[i] for i in indices[0] if 0 <= i < len(ids)]
        except Exception as e:
            logger.error(f"FAISS search error: {e}")
            return []
